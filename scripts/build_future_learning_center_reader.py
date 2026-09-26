from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


THEMES = {
    "T1": "Future Learning Center Concepts and Theories",
    "T2": "Library Space Transformation and Learning Commons",
    "T3": "Learning Support Services",
    "T4": "User Needs, Experience and Behaviour",
    "T5": "Digital Intelligence, AI and Smart Library Services",
    "T6": "Organizational Governance and Service Models",
    "T7": "International Comparison and Foreign Practices",
}

UNRELIABLE_AUTHORS = {
    "FLC-0A78AE1EC059",
    "FLC-E86EDD77A7D8",
    "FLC-BFA71761FFA6",
    "FLC-B21262428AA9",
    "FLC-D0D577A4BEE2",
    "FLC-5A559DBBB63F",
    "FLC-54FCF9633BD1",
    "FLC-740E06115005",
    "FLC-AAF1547D60BA",
    "FLC-AF43D204976D",
}

KEY_NOTES = [
    {
        "id": "FLC-D2B84F8B5474",
        "question": "高校未来学习中心的智能服务需求应如何分类和排序，才能避免技术炫技并回应真实教学需求？",
        "contribution": "把智能服务需求区分为必备型、期望型、魅力型和无差异型，并给出分层配置逻辑。",
        "limitation": "论文明确报告局限，但自动提取的局限文本不完整；其问卷以在校学生为主，跨校、跨群体外推仍需谨慎。",
        "why": "它把“建设什么”转化为可调查、可排序的需求变量，是用户驱动研究的清晰起点。",
    },
    {
        "id": "FLC-0A78AE1EC059",
        "question": "高校未来学习中心应提供哪些服务内容，不同需求属性应如何转化为优化优先级？",
        "contribution": "报告3项必备型、9项期望型、6项魅力型和8项无差异型需求，并从六类服务维度提出优化建议。",
        "limitation": "论文明确承认不足，但局限段落的文本提取受版面串行影响；正式问卷的覆盖和代表性需回到原文核查。",
        "why": "它连接了用户需求分类与具体服务组合，适合用于设计后续问卷或服务评价框架。",
    },
    {
        "id": "FLC-1C677BADE9E4",
        "question": "国内外高校图书馆空间服务在未来学习中心语境下有哪些可比较的差异？",
        "contribution": "以国内外各20所高校图书馆为调查对象，将空间开放、预约和服务呈现方式放进同一比较框架。",
        "limitation": "研究问题在数据库中为inferred，论文未明确报告limitations；网站呈现不等同于实际使用效果。",
        "why": "它为跨机构空间服务比较提供了可复用的观察维度，也暴露了“网页可见性”与“真实体验”之间的测量差距。",
    },
    {
        "id": "FLC-5A559DBBB63F",
        "question": "ARL成员高校图书馆与我国“双一流”高校图书馆的空间服务差异能提供哪些建设启示？",
        "contribution": "把ARL成员馆与147所国内“双一流”高校图书馆纳入比较，扩大了国际比较的机构覆盖。",
        "limitation": "论文明确报告不足；调查聚焦高校图书馆网站和内容分析，制度环境差异及线下运行成效仍难控制。",
        "why": "它是少数把国际案例从单馆叙述推进到同口径机构比较的论文。",
    },
    {
        "id": "FLC-E86EDD77A7D8",
        "question": "应用型本科院校的师生需求是否要求不同于研究型高校的未来学习中心服务组合？",
        "contribution": "以广州应用科技学院为场域，使用KANO与Better-Worse分析，将应用型人才培养需求转化为服务优先级。",
        "limitation": "论文明确指出样本受限于特定院校与抽样方式，结论不宜直接推广到其他类型高校。",
        "why": "它提醒研究者不能把“双一流”或研究型大学经验视为所有高校的默认模型。",
    },
    {
        "id": "FLC-54FCF9633BD1",
        "question": "高校图书馆员向“学习策划者”转型需要承担什么职责、具备什么能力？",
        "contribution": "将学习资源整合、数据分析支持、跨学科学习和教学设计咨询连接到馆员角色转型，并报告馆员与师生调查。",
        "limitation": "论文提出需要继续研究学习策划的长效机制与实践推广；现有证据更接近需求与态度，而非长期成效。",
        "why": "它把未来学习中心从设施项目转向人员能力和服务劳动，是组织研究的重要入口。",
    },
    {
        "id": "FLC-740E06115005",
        "question": "面向未来学习中心的高校图书馆智慧空间研究形成了哪些热点与方向？",
        "contribution": "用文献、内容和计量分析识别服务创新、用户体验优化及数智技术应用三个热点。",
        "limitation": "数据库未识别到明确limitations；文献热点并不等于建设成效，也不能代替用户层面的效果证据。",
        "why": "它适合作为智慧空间专题的入口文献，并可用于定位后续经验研究应补足的位置。",
    },
    {
        "id": "FLC-AAF1547D60BA",
        "question": "用户需求与行为特征如何影响未来学习中心背景下的文献资源建设？",
        "contribution": "报告1,031份有效问卷，覆盖石河子大学19个学院及行政教辅单位，把资源建设与大样本用户数据连接起来。",
        "limitation": "数据库未识别到明确limitations；单校调查的组织结构和学科分布可能限制外部效度。",
        "why": "它提供了少见的大样本资源需求证据，有助于抵消“空间和技术即转型”的单线叙事。",
    },
    {
        "id": "FLC-B21262428AA9",
        "question": "高校学习空间研究如何解释空间形态与学习成效之间的关系？",
        "contribution": "结合文献统计和建构主义学习理论，强调以学生为中心、支持协作的空间与认知及非认知能力发展之间的关系。",
        "limitation": "数据库未识别到明确limitations；其样本字段来自所综述文献的研究对象，不能视为一项统一原始样本。",
        "why": "它帮助研究者从“空间配置”转向“空间如何支持学习机制”的问题。",
    },
    {
        "id": "FLC-AF0DA0F0CDAE",
        "question": "新文科情境下，高校图书馆未来学习中心应包含哪些设计要素？",
        "contribution": "使用扎根理论，从案例资料中组织空间、技术、学习支持、社会活动和学习情境等要素。",
        "limitation": "数据库未识别到明确limitations；案例资料来自公开渠道，资料可见性可能影响范畴饱和度。",
        "why": "它提供了学科情境化的设计模型，可与通用建设清单形成对照。",
    },
    {
        "id": "FLC-F5599DA2A77B",
        "question": "未来学习中心通过哪些要素和路径赋能高校图书馆高质量发展？",
        "contribution": "以扎根理论归纳政策、管理、空间、资源、技术、服务、育人和教育八类要素。",
        "limitation": "数据库未识别到明确limitations；材料样本为筛选后的未来学习中心文献，模型仍需在实际机构中验证。",
        "why": "它是把分散建设要素组织为中层模型的代表，可用于形成可检验命题。",
    },
    {
        "id": "FLC-77F19452C268",
        "question": "我国未来学习中心研究的主题如何识别，并经历了怎样的演化？",
        "contribution": "通过信息计量与内容分析，将研究概括为理论体系、建设路径和发展创新实践三条主题线。",
        "limitation": "论文明确提出未来仍需解决知识单元动态标引与跨平台资源协同等问题；计量结果受数据库和检索式影响。",
        "why": "它提供了本语料之外的领域演化参照，适合用于检验本数据库主题图是否过度依赖当前语料。",
    },
    {
        "id": "FLC-AF43D204976D",
        "question": "未来学习空间如何影响教学过程、学生参与和学习体验？",
        "contribution": "以北京师范大学未来学习体验中心为案例，结合问卷与访谈评价空间应用效果。",
        "limitation": "论文明确指出中心仍在优化，硬件不足可能影响满意度；个案情境限制结论推广。",
        "why": "它是语料中较早、也较少见的使用后效果评价，为后续纵向或准实验研究提供基线。",
    },
    {
        "id": "FLC-9DD904535079",
        "question": "哪些因素影响高校图书馆用户采纳未来学习中心？",
        "contribution": "以质性研究识别认知、技术和环境三类影响因素，并报告编码一致性系数0.916。",
        "limitation": "数据库未识别到明确limitations，且出版年未可靠提取；质性因素模型仍需量化检验和跨场域验证。",
        "why": "它把研究对象从建设供给转到用户采纳过程，可连接TAM、UTAUT等成熟采纳研究传统。",
    },
    {
        "id": "FLC-FADDE187EAA5",
        "question": "LLaMA大模型嵌入高校未来学习中心会带来哪些法律与治理风险？",
        "contribution": "识别内容可靠性、隐私泄露和著作权权属三类风险，并提出技术、制度、法律和协同治理路径。",
        "limitation": "数据库未识别到明确limitations；方法描述为技术解构与场景验证，但未在methodology字段中形成可复核的样本信息。",
        "why": "它为AI赋能叙事提供风险和责任边界，是技术乐观论的重要对照。",
    },
    {
        "id": "FLC-079C3D09D5B5",
        "question": "高校图书馆如何通过未来学习中心数智赋能教学科研服务？",
        "contribution": "提出融合空间、智能氛围、交互支持、万物智联、创新能力和国际视野六个方面的6I框架。",
        "limitation": "数据库未识别到明确methodology、findings或limitations；框架主要用于组织建设议题，仍需经验检验。",
        "why": "它提供了高度可读的综合框架，也适合作为检验“框架是否可测量、可落地”的对象。",
    },
]


def read_records(root: Path) -> list[dict]:
    path = root / "processed" / "manifest.jsonl"
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line]


def ids(*values: str) -> str:
    return " ".join(f"[{value}]" for value in values)


def field_label(item: dict) -> str:
    pages = ",".join(map(str, item["evidence_pages"]))
    return f"{item['status']}" + (f", evidence pages {pages}" if pages else "")


def clean_author(record: dict) -> str:
    author = record["metadata"].get("author", "").strip()
    if not author or record["stable_id"] in UNRELIABLE_AUTHORS:
        return "作者元数据未可靠提取"
    return author


def citation(record: dict) -> str:
    metadata = record["metadata"]
    year = metadata.get("year") or "年份未可靠提取"
    journal = metadata.get("journal") or "期刊元数据未可靠提取"
    return f"{clean_author(record)}：《{metadata['title']}》，{journal}，{year}。[{record['stable_id']}]"


def build_reader(records: list[dict]) -> str:
    by_id = {record["stable_id"]: record for record in records}
    included = [r for r in records if r["theme"]["scope_status"] == "included"]
    years = Counter(r["metadata"]["year"] for r in included if r["metadata"]["year"])
    journals = Counter(r["metadata"]["journal"] for r in included if r["metadata"]["journal"])
    primary = Counter(r["theme"]["primary_theme"] for r in included)
    methods = Counter(method for r in included for method in r["analysis"]["methods"])
    theme_methods = {
        theme: Counter(
            method
            for r in included
            if r["theme"]["primary_theme"] == theme
            for method in r["analysis"]["methods"]
        )
        for theme in THEMES
    }
    explicit_methods = sum(r["analysis"]["methodology"]["status"] == "explicit" for r in records)
    explicit_samples = sum(r["analysis"]["sample"]["status"] == "explicit" for r in records)
    explicit_findings = sum(r["analysis"]["findings"]["status"] == "explicit" for r in records)
    explicit_limits = sum(r["analysis"]["limitations"]["status"] == "explicit" for r in records)

    year_rows = "\n".join(
        f"| {year} | {count} | {ids(next(r['stable_id'] for r in included if r['metadata']['year'] == year))} |"
        for year, count in sorted(years.items())
    )
    journal_rows = "\n".join(
        f"| {journal} | {count} | {ids(next(r['stable_id'] for r in included if r['metadata']['journal'] == journal))} |"
        for journal, count in journals.most_common(10)
    )

    theme_specs = [
        (
            "T1",
            "把未来学习中心作为概念、组织形态、要素系统或理论对象，讨论其边界、构成和作用机制。",
            ("FLC-BFA71761FFA6", "FLC-A8C1119CD0C8", "FLC-08C45D85DED5", "FLC-77F19452C268"),
            "未来学习中心究竟是空间、服务平台、基层学习组织，还是图书馆转型的政策性标签？其关键要素如何关联？",
            "少量论文使用Bibliometric analysis或Grounded theory，更多论文未明确报告方法。",
            "共同点是强调人、空间、资源、技术、服务和治理的组合；差异在于是否把“学习组织”放在中心，还是把图书馆建设要素放在中心。",
            "核心概念尚缺少稳定的操作化定义；同一术语可能指组织、空间、项目或服务组合。",
        ),
        (
            "T2",
            "关注图书馆物理与虚拟空间、Learning Commons、智慧空间及其支持学习的机制。",
            ("FLC-1C677BADE9E4", "FLC-B21262428AA9", "FLC-AF43D204976D", "FLC-6A2498CAFD73"),
            "空间如何配置、开放、预约和治理？空间是否促进协作、参与、满意度或学习成效？",
            "本主题较多使用Literature review、Web survey、Bibliometric analysis，少量采用问卷、访谈和案例评价。",
            "研究普遍支持从藏书空间转向复合、共享、协作和虚实融合空间，但只有少数论文直接评价使用效果。",
            "空间特征与学习结果之间的因果链条仍不清楚；网页调查、设计说明和真实使用行为不可互换。",
        ),
        (
            "T3",
            "研究教学支持、科研支持、学业规划、知识服务和学习策划等面向学习过程的服务。",
            ("FLC-079C3D09D5B5", "FLC-4979619F5099", "FLC-54FCF9633BD1", "FLC-93ABC817D493"),
            "图书馆能提供哪些学习支持？服务如何嵌入课程、科研和学生发展？馆员需要怎样的新角色？",
            "该主题仅少量论文明确报告Web survey等方法，方法透明度低于用户需求主题。",
            "论文通常主张从资源供给走向持续的教学科研协同、个性化支持和学习策划。",
            "服务使用、学习结果、馆员工作负荷和跨部门责任边界缺少持续评估。",
        ),
        (
            "T4",
            "以用户需求、体验、行为和采纳为中心，把建设内容转化为可调查的偏好、影响因素或服务优先级。",
            ("FLC-D2B84F8B5474", "FLC-0A78AE1EC059", "FLC-40CAC33A5980", "FLC-9DD904535079"),
            "用户真正需要什么？哪些因素影响满意度、采纳和持续使用？不同群体是否需要不同服务组合？",
            "这是经验方法最集中的主题：Questionnaire survey、KANO analysis、Interview、Grounded theory和Web survey均有出现。",
            "KANO研究表明需求有不同属性，不应将所有功能视为同等优先；体验研究则提示环境、技术、服务和学习体验需同时考虑。",
            "单校样本、便利抽样和横截面设计较常见；用户态度尚未充分连接到长期学习行为和成效。",
        ),
        (
            "T5",
            "讨论AI、AIGC、大语言模型、智能体、数字孪生、智慧服务及相应风险治理。",
            ("FLC-8177904692F9", "FLC-FADDE187EAA5", "FLC-282F29E53076", "FLC-C4F983ADAB67"),
            "新技术可以嵌入哪些学习支持场景？如何处理准确性、隐私、版权、偏差和人机责任？",
            "多数论文是概念或场景讨论，少量使用案例、文献计量和内容分析；可复核的服务效果评价有限。",
            "技术论文强调个性化、主动服务和资源整合，同时风险论文强调数据可靠性、隐私和权属治理。",
            "“可实现”常被写成“有效”；缺少与传统服务的对照、失败案例、成本和责任分配研究。",
        ),
        (
            "T6",
            "关注顶层设计、跨部门协同、组织治理、服务模式、馆员能力和资源整合。",
            ("FLC-5219CF870983", "FLC-CFCDAA38CDA0", "FLC-54FCF9633BD1", "FLC-1FF3D56945F1"),
            "由谁牵头、如何协同、怎样配置资源和能力？未来学习中心是图书馆项目还是校级学习基础设施？",
            "本主题方法较分散，包括Literature review、Web survey、Case study和少量结构方程、访谈或扎根理论。",
            "多数论文认为仅改造空间不足，必须同时处理政策、组织、人员、资源和服务流程。",
            "治理建议多而实施证据少；权责、预算、绩效和跨部门冲突尚未形成可比较指标。",
        ),
        (
            "T7",
            "使用国外单馆案例、跨国案例或国际机构比较，为国内实践提供参照。",
            ("FLC-5A559DBBB63F", "FLC-1FBEE4B5422D", "FLC-B5EF8C13E992", "FLC-053F5A168B8B"),
            "国外学习中心如何组织空间、教学与技术？哪些做法可迁移，哪些依赖制度和文化情境？",
            "主题规模很小，明确方法主要是Literature review与Content analysis；单案例介绍仍占重要位置。",
            "国外案例常被用于证明互动空间、数字课堂和教学协作的可能性。",
            "制度、经费、用户结构和治理环境通常未被同口径控制，容易把案例启示误当作可直接复制的模式。",
        ),
    ]

    theme_sections = []
    for theme, definition, reps, questions, typical, conclusions, unresolved in theme_specs:
        method_text = "；".join(f"{name} {count}" for name, count in theme_methods[theme].most_common(4)) or "无稳定的明确方法分布"
        theme_sections.append(
            f"""### {THEMES[theme]}（{theme}）

**定义。** {definition} 该主题在纳入语料中有{primary[theme]}篇主主题论文。{ids(*reps)}

**常见研究问题。** {questions} {ids(*reps[:2])}

**典型方法。** {typical} 数据库中明确出现频次较高的方法为：{method_text}；方法计数可重叠，且未报告方法的论文不被推定为“文献研究法”。{ids(*reps)}

**主要结论。** {conclusions} 这里的“结论”是对该主题论文论证方向的归纳，不表示所有论文共享同一证据强度。{ids(*reps)}

**未解决问题。** {unresolved} {ids(*reps)}
"""
        )

    key_reader_ids = [
        "FLC-D2B84F8B5474",
        "FLC-0A78AE1EC059",
        "FLC-1C677BADE9E4",
        "FLC-5A559DBBB63F",
        "FLC-54FCF9633BD1",
        "FLC-B21262428AA9",
        "FLC-AF43D204976D",
        "FLC-77F19452C268",
        "FLC-9DD904535079",
        "FLC-FADDE187EAA5",
        "FLC-079C3D09D5B5",
    ]
    note_map = {note["id"]: note for note in KEY_NOTES}
    key_sections = []
    for stable_id in key_reader_ids:
        record = by_id[stable_id]
        note = note_map[stable_id]
        methodology = record["analysis"]["methodology"]
        key_sections.append(
            f"""### {record['metadata']['title']} [{stable_id}]

**Citation information.** {citation(record)}

**Research question.** {note['question']} 数据库状态：{field_label(record['analysis']['research_question'])}。

**Method.** {methodology['value'] or 'not_reported'}（{field_label(methodology)}）。[{stable_id}]

**Contribution.** {note['contribution']} [{stable_id}]

**Limitation.** {note['limitation']} [{stable_id}]

**Why this paper matters.** {note['why']} [{stable_id}]
"""
        )

    return f"""# Future Learning Center Literature Reader

> **Reader purpose.** 这是一份研究笔记型读本，而不是一篇已经完成的学术综述。它帮助研究者识别概念、证据、争议与可行问题；不预设最终选题，也不把所有论文强行解释为同一条“未来学习中心”发展路线。所有方括号中的FLC编号均为数据库stable_id。

## 1. Corpus Overview

### 1.1 Corpus boundary and evidence convention

语料共有215份PDF，214份完成文本提取，1份零字节文件保留失败状态；210篇进入主题分析，5篇标记为background_only。下文的计数是对ID级记录的数据库汇总，而论文层面的主张均附具体ID。失败记录是[FLC-E3B0C44298FC]，没有被静默删除。

数据库区分`explicit`、`inferred`和`not_reported`。`inferred`通常表示研究问题从题名或明确论述归纳而来；`not_reported`表示没有足够证据写入该字段，不等于作者绝对没有使用该方法或讨论该问题。这个边界对阅读概念性和规范性论文尤其重要。{ids("FLC-BFA71761FFA6", "FLC-079C3D09D5B5", "FLC-D2B84F8B5474")}

### 1.2 Publication years

在有可靠年份的纳入论文中，2024年57篇、2025年82篇，显示语料明显集中于近期；另有24篇纳入论文的年份未可靠提取。因此，不宜把这套语料当成2013年以来均匀抽样的长期序列。{ids("FLC-079C3D09D5B5", "FLC-D2B84F8B5474", "FLC-77F19452C268")}

| Year | Papers with reliable year | Example ID |
|---:|---:|---|
{year_rows}

### 1.3 Major journals

在期刊字段可靠的纳入论文中，发文较多的来源包括《图书馆建设》《图书馆学刊》《江苏科技信息》《图书情报工作》和《大学图书情报学刊》。共有58条记录的journal字段缺失，所以这些数字描述的是已识别来源，而不是完整的期刊份额。{ids("FLC-39D9F87A8F5B", "FLC-D2B84F8B5474", "FLC-0A78AE1EC059", "FLC-F04D86A4AE33")}

| Journal | Identified papers | Example ID |
|---|---:|---|
{journal_rows}

### 1.4 Research characteristics

第一，语料是“建设导向”和“问题导向”的混合体。相当多论文讨论概念、路径、场景和治理，而用户需求、空间效果和采纳研究形成了较小但更经验化的证据群。{ids("FLC-BFA71761FFA6", "FLC-CFCDAA38CDA0", "FLC-D2B84F8B5474", "FLC-AF43D204976D")}

第二，主题之间高度交叉。空间论文经常涉及服务，AI论文涉及治理，用户研究又反过来影响空间与技术优先级；primary_theme只是阅读入口，不是互斥的学科边界。{ids("FLC-6A2498CAFD73", "FLC-C4F983ADAB67", "FLC-40CAC33A5980")}

第三，证据透明度有限：{explicit_methods}篇明确报告methodology，{explicit_samples}篇明确报告sample，{explicit_findings}篇明确报告findings，只有{explicit_limits}篇明确报告limitations。因而，读者应把“模型/路径建议”和“经验证据支持的结论”分开处理。{ids("FLC-D2B84F8B5474", "FLC-AF43D204976D", "FLC-079C3D09D5B5")}

## 2. Conceptual Background

### 2.1 What “future learning center” means in this corpus

本语料没有唯一、稳定的定义。至少存在三种用法。其一，把未来学习中心视为依托高校图书馆的新型基层学习组织，核心是主体、资源和技术要素的组合。{ids("FLC-BFA71761FFA6")}

其二，把它视为高校图书馆数智赋能教学科研、重组空间与服务的转型平台，图书馆可以牵头但需要校内外协作。{ids("FLC-079C3D09D5B5", "FLC-CFCDAA38CDA0")}

其三，把它视为面向未来学习的综合生态：物理空间、数字平台、学习活动、馆员能力和治理机制共同工作。这个解释更接近“系统”而非“场所”。{ids("FLC-4979619F5099", "FLC-54FCF9633BD1", "FLC-6A2498CAFD73")}

> **Interpretive caution.** 这些用法可以互补，也可能竞争。若研究不先说明分析单位是组织、空间、服务系统还是政策项目，同一“建设成效”指标可能测量完全不同的对象。{ids("FLC-BFA71761FFA6", "FLC-AF43D204976D", "FLC-CFCDAA38CDA0")}

### 2.2 Related concepts and distinctions

| Concept | Primary unit | Central question | Relationship in this corpus |
|---|---|---|---|
| Learning center | 组织或服务单元 | 如何集中支持学习、教学与资源使用？ | 是较宽的功能概念，不必然以高技术或图书馆转型为前提。[FLC-EF061AD82805] |
| Learning commons | 空间—资源—服务组合 | 如何通过共享、协作和灵活空间支持学习？ | 更强调空间与服务的耦合，可成为未来学习中心的一部分，但不等同于完整治理系统。[FLC-50A81720A069] [FLC-B21262428AA9] |
| Smart library | 技术与基础设施体系 | 如何用感知、数据和自动化改善资源与服务？ | 与未来学习中心有“内涵耦合”，但技术智能不自动产生学习支持或学习成效。[FLC-39D9F87A8F5B] [FLC-C4F983ADAB67] |
| Library transformation | 组织变化过程 | 图书馆的空间、角色、流程和治理如何改变？ | 是更宽的过程概念；未来学习中心只是其中一种政策与实践载体。[FLC-1FF3D56945F1] [FLC-CFCDAA38CDA0] |

概念差异最重要的不是名称，而是研究对象和结果变量。若把Learning Commons当作空间，结果变量可能是使用、协作和满意度；若把未来学习中心当作组织，结果变量可能是跨部门协同、服务覆盖和学习支持；若把Smart Library当作技术体系，则还要测量准确性、隐私、采用和维护成本。{ids("FLC-AF43D204976D", "FLC-54FCF9633BD1", "FLC-FADDE187EAA5")}

## 3. Research Theme Map

主题数量分别为：T1 {primary["T1"]}篇、T2 {primary["T2"]}篇、T3 {primary["T3"]}篇、T4 {primary["T4"]}篇、T5 {primary["T5"]}篇、T6 {primary["T6"]}篇、T7 {primary["T7"]}篇。计数仅表示primary_theme；同一论文还可有secondary_themes。{ids("FLC-BFA71761FFA6", "FLC-1C677BADE9E4", "FLC-D2B84F8B5474", "FLC-FADDE187EAA5", "FLC-CFCDAA38CDA0")}

{"".join(theme_sections)}

## 4. Methodological Landscape

### 4.1 Survey and user studies

明确方法中，Questionnaire survey出现{methods["Questionnaire survey"]}篇，Web survey出现{methods["Web survey"]}篇，Interview出现{methods["Interview"]}篇，KANO analysis出现{methods["KANO analysis"]}篇。方法计数可重叠。KANO研究能把需求转化为优先级，问卷和访谈能呈现用户差异；但单校横截面、便利抽样和态度指标限制了对长期使用与学习结果的解释。{ids("FLC-D2B84F8B5474", "FLC-E86EDD77A7D8", "FLC-A1CA30499358", "FLC-AF43D204976D")}

### 4.2 Case and comparative studies

Case study明确出现{methods["Case study"]}篇，另有多篇以单馆实践或网站调查进行比较。案例研究擅长解释建设过程和情境，跨馆网站调查擅长建立可比清单；两者的共同限制是“公开呈现”可能不同于“实际运行”，成功案例也可能产生选择偏差。{ids("FLC-46F65852D8D1", "FLC-1FBEE4B5422D", "FLC-1C677BADE9E4", "FLC-5A559DBBB63F")}

### 4.3 Theoretical and model-building discussions

明确识别的理论或模型包括Constructivist Learning Theory、KANO Model、TOE Framework、Supply-Demand Theory、Activity Theory、UTAUT和Technology Acceptance Model等，但绝大多数论文没有明确理论框架。理论论文的价值在于提供概念关系和解释机制，其风险是把建设要素清单误当作已经验证的因果模型。{ids("FLC-0970AAAB7995", "FLC-5219CF870983", "FLC-08C45D85DED5", "FLC-A8C1119CD0C8", "FLC-CC494CC9ADA4")}

### 4.4 Empirical evaluation

真正把建设投入与结果联系起来的评价研究较少。北京师范大学未来学习体验中心研究结合问卷与访谈讨论教学过程、学生参与和满意度；用户体验、采纳意愿和空间价值研究则开始构造影响因素或评价维度。它们说明“效果评价”可行，但目前仍缺少纵向、对照和跨校设计。{ids("FLC-AF43D204976D", "FLC-E1B3B8EC4827", "FLC-40CAC33A5980", "FLC-9DD904535079")}

> **Method note.** 方法未报告不等于没有方法；但在证据综合中，不能用题名或一般论述替作者补写方法。任何准备引用为“实证证明”的论文，都应再次核对其methodology、sample和evidence_pages。{ids("FLC-079C3D09D5B5", "FLC-BFA71761FFA6", "FLC-D2B84F8B5474")}

## 5. Key Papers

本节以priority_reading_list为起点，并补入早期效果评价、用户采纳、综合框架和AI风险治理论文，以形成互补阅读路径。排序不是引用次数，也不是最终权威名单。

{"".join(key_sections)}

## 6. Research Debates

### 6.1 Technology-driven vs user-driven transformation

技术驱动论文把AI、大模型和智能体视为主动服务、个性化学习和资源整合的基础；用户驱动论文则指出需求有必备、期望、魅力和无差异之分，技术功能必须接受用户价值排序。二者并非直接互相反驳，但对“从哪里开始建设”给出不同答案。{ids("FLC-8177904692F9", "FLC-C4F983ADAB67", "FLC-D2B84F8B5474", "FLC-0A78AE1EC059")}

### 6.2 Physical space vs service innovation

空间研究把灵活、共享、沉浸和虚实融合视为学习方式变化的条件；学习支持研究则强调空间只有与教学、科研、资源和馆员服务相连才有意义。争议不在“要不要空间”，而在空间是结果本身，还是服务和学习活动的基础设施。{ids("FLC-B21262428AA9", "FLC-1C677BADE9E4", "FLC-4979619F5099", "FLC-54FCF9633BD1")}

### 6.3 Institutional planning vs empirical user research

治理与规划论文通常从政策、资源、组织和路径出发，强调顶层设计和跨部门协同；用户研究从需求、体验和采纳出发，要求建设优先级能被调查和验证。前者擅长解释“如何组织”，后者擅长检验“对谁有用”。{ids("FLC-CFCDAA38CDA0", "FLC-5219CF870983", "FLC-E86EDD77A7D8", "FLC-9DD904535079")}

### 6.4 Library-led vs institution-wide ownership

部分论文主张高校图书馆应成为牵头机构，以资源、空间和专业能力连接校内部门；另一些治理框架实际上把未来学习中心描述为需要教学管理、院系、技术部门和外部平台共同参与的校级系统。这里存在领导权、责任和预算归属的问题，而不是简单的概念差异。{ids("FLC-079C3D09D5B5", "FLC-CFCDAA38CDA0", "FLC-5219CF870983")}

### 6.5 Innovation optimism vs risk governance

AI应用论文强调智能体和大模型的服务潜力，风险研究则把内容可靠性、隐私、版权和人机责任置于同等位置。语料尚未形成统一治理标准，因此“技术可用”与“组织可接受”必须分开判断。{ids("FLC-8177904692F9", "FLC-7E32D0AFAA75", "FLC-FADDE187EAA5", "FLC-282F29E53076")}

## 7. Research Gaps

### 7.1 Insufficiently studied topics

学习成效、持续使用、服务可达性、公平性、特殊群体和成本效益研究不足。现有研究更常测量建设要素、需求或满意度，而不是长期学习结果。{ids("FLC-AF43D204976D", "FLC-40CAC33A5980", "FLC-A1CA30499358")}

非“双一流”、职业院校、应用型本科和学科特定场景虽已出现，但仍未构成可比较的类型学。单个情境的发现不能自动推广到所有高校。{ids("FLC-E86EDD77A7D8", "FLC-309D35433AE0", "FLC-0970AAAB7995", "FLC-AF0DA0F0CDAE")}

### 7.2 Methodological limitations

160篇论文未明确报告methodology，167篇未明确报告sample，159篇未明确报告findings，180篇未明确报告limitations。即使考虑自动提取误差，这仍提示研究透明度是重要问题。{ids("FLC-D2B84F8B5474", "FLC-AF43D204976D", "FLC-079C3D09D5B5")}

横截面问卷、公开网站调查和成功案例较多，纵向研究、对照设计、跨校同口径数据和行为日志较少。由此得到的多数结论更适合描述、分类和提出命题，而不适合做因果判断。{ids("FLC-1C677BADE9E4", "FLC-5A559DBBB63F", "FLC-D2B84F8B5474", "FLC-AF43D204976D")}

### 7.3 Theoretical gaps

明确理论框架集中在少数论文，且不同理论常各自出现一次，尚未形成持续累积的变量体系。概念模型、采纳模型、空间机制和治理框架之间缺乏可比较的共同结果变量。{ids("FLC-A8C1119CD0C8", "FLC-5219CF870983", "FLC-CC494CC9ADA4", "FLC-08C45D85DED5")}

“未来学习中心”的分析单位仍不稳定：组织、空间、技术系统和服务组合经常共用同一名称。理论推进需要先界定对象，再说明机制和结果，而不是继续扩展建设要素清单。{ids("FLC-BFA71761FFA6", "FLC-39D9F87A8F5B", "FLC-CFCDAA38CDA0")}

## 8. Possible Research Directions

以下方向互不构成优先级，也不预设最终选题。选择时应依据研究者可进入的场域、可获得的数据、时间跨度和理论兴趣。

### Path A: Post-occupancy learning-space evaluation

**Research question.** 空间改造后，学生参与、协作、持续使用和学习结果是否发生变化？不同空间要素通过什么机制产生影响？{ids("FLC-AF43D204976D", "FLC-B21262428AA9", "FLC-40CAC33A5980")}

**Potential contribution.** 把空间设计研究从建设说明推进到结果评价，并区分满意度、行为和学习成效。{ids("FLC-AF43D204976D", "FLC-B21262428AA9", "FLC-40CAC33A5980")}

**Possible difficulty.** 需要改造前基线、对照空间或长期追踪；课程、教师和学生自选择会干扰因果解释。{ids("FLC-AF43D204976D")}

### Path B: User-needs portfolio and service prioritization

**Research question.** 不同高校类型和用户群体的必备、期望与魅力需求是否稳定？需求排序如何随使用经验变化？{ids("FLC-D2B84F8B5474", "FLC-0A78AE1EC059", "FLC-E86EDD77A7D8")}

**Potential contribution.** 建立可比较的服务优先级工具，检验KANO结果的跨校稳定性与时间变化。{ids("FLC-D2B84F8B5474", "FLC-0A78AE1EC059", "FLC-E86EDD77A7D8")}

**Possible difficulty.** 问卷项目需要等值性检验；用户表达需求不等于实际使用，需结合行为数据。{ids("FLC-D2B84F8B5474", "FLC-E86EDD77A7D8")}

### Path C: AI-agent service effectiveness and governance

**Research question.** AI-Agent在咨询、学习策划或资源推荐中是否优于现有服务？准确性、隐私、版权和馆员责任如何共同评价？{ids("FLC-7E32D0AFAA75", "FLC-8177904692F9", "FLC-FADDE187EAA5")}

**Potential contribution.** 将功能设想转化为服务实验和治理指标，连接技术性能、用户结果和制度责任。{ids("FLC-7E32D0AFAA75", "FLC-FADDE187EAA5")}

**Possible difficulty.** 模型快速迭代、数据授权和风险暴露使长期可重复性与伦理审批更复杂。{ids("FLC-FADDE187EAA5", "FLC-282F29E53076")}

### Path D: Governance and cross-department collaboration

**Research question.** 图书馆牵头、联合治理和校级平台三种模式如何影响资源配置、服务整合和责任承担？{ids("FLC-079C3D09D5B5", "FLC-5219CF870983", "FLC-CFCDAA38CDA0")}

**Potential contribution.** 把“顶层设计”拆成可观察的决策权、预算、流程、绩效和冲突解决机制。{ids("FLC-5219CF870983", "FLC-CFCDAA38CDA0")}

**Possible difficulty.** 组织数据敏感，案例数量有限，且高校治理结构差异会降低直接可比性。{ids("FLC-079C3D09D5B5", "FLC-CFCDAA38CDA0")}

### Path E: Learning-support service and librarian role

**Research question.** 学习策划、教学设计咨询和数据分析支持如何嵌入课程？馆员能力与服务结果之间是什么关系？{ids("FLC-54FCF9633BD1", "FLC-4979619F5099", "FLC-079C3D09D5B5")}

**Potential contribution.** 把馆员角色讨论转化为任务、能力、工作负荷和学习支持结果的模型。{ids("FLC-54FCF9633BD1", "FLC-4979619F5099")}

**Possible difficulty.** 服务往往由多部门共同提供，难以单独识别图书馆贡献；需要获得课程和服务过程数据。{ids("FLC-54FCF9633BD1", "FLC-079C3D09D5B5")}

### Path F: Cross-institutional and international comparison

**Research question.** 在统一指标下，不同高校类型或国家的空间、服务和治理模式有哪些稳定差异？{ids("FLC-1C677BADE9E4", "FLC-5A559DBBB63F", "FLC-1FBEE4B5422D")}

**Potential contribution.** 区分可迁移机制与情境依赖做法，减少以单个成功案例代替比较证据的问题。{ids("FLC-5A559DBBB63F", "FLC-1FBEE4B5422D")}

**Possible difficulty.** 公开信息口径不一致，制度、经费和用户结构差异需要显式控制。{ids("FLC-1C677BADE9E4", "FLC-5A559DBBB63F")}

### Path G: Concept and measurement clarification

**Research question.** 未来学习中心作为组织、空间、服务系统和政策项目时，各自的必要条件、边界和结果变量是什么？{ids("FLC-BFA71761FFA6", "FLC-39D9F87A8F5B", "FLC-77F19452C268")}

**Potential contribution.** 建立概念矩阵和测量模型，为后续实证研究提供共同语言。{ids("FLC-BFA71761FFA6", "FLC-77F19452C268")}

**Possible difficulty.** 过度追求统一定义可能抹去不同院校的合理差异；需要保留多层次概念结构。{ids("FLC-BFA71761FFA6", "FLC-39D9F87A8F5B")}

### Path H: Institution-type and equity differences

**Research question.** 应用型本科、高职、中职、医学和新文科等场景的需求、资源约束和成效标准有何差异？{ids("FLC-E86EDD77A7D8", "FLC-309D35433AE0", "FLC-0970AAAB7995", "FLC-AF0DA0F0CDAE")}

**Potential contribution.** 形成高校类型与用户群体的情境化模型，避免把资源充足高校经验作为默认标准。{ids("FLC-E86EDD77A7D8", "FLC-309D35433AE0")}

**Possible difficulty.** 机构类型内部差异很大，样本获取和指标等值性会成为研究设计难点。{ids("FLC-E86EDD77A7D8", "FLC-AF0DA0F0CDAE")}

> **Decision note.** 这些路径可组合，但不宜在一个研究中同时承担概念澄清、平台建设、用户调查、效果评价和治理设计。较稳妥的下一步是先确定分析单位与可获得证据，再决定问题，而不是先选择“未来学习中心”作为总题目。{ids("FLC-BFA71761FFA6", "FLC-D2B84F8B5474", "FLC-CFCDAA38CDA0")}
"""


def build_key_notes(records: list[dict]) -> str:
    by_id = {record["stable_id"]: record for record in records}
    sections = []
    for index, note in enumerate(KEY_NOTES, 1):
        record = by_id[note["id"]]
        analysis = record["analysis"]
        secondary = [
            theme
            for theme in record["theme"]["secondary_themes"]
            if theme != record["theme"]["primary_theme"]
        ]
        sections.append(
            f"""## {index}. {record['metadata']['title']} [{note['id']}]

**Citation information.** {citation(record)}

**Primary theme.** {THEMES[record['theme']['primary_theme']]} ({record['theme']['primary_theme']}); secondary themes: {", ".join(secondary) or "none recorded"}. [{note['id']}]

**Research question.** {note['question']} Database status: {field_label(analysis['research_question'])}. [{note['id']}]

**Method.** {analysis['methodology']['value'] or "not_reported"} ({field_label(analysis['methodology'])}). Sample status: {field_label(analysis['sample'])}. [{note['id']}]

**Contribution.** {note['contribution']} [{note['id']}]

**Limitation.** {note['limitation']} Database limitation status: {field_label(analysis['limitations'])}. [{note['id']}]

**Why this paper matters.** {note['why']} [{note['id']}]
"""
        )
    return f"""# Key Paper Notes

These notes start from the validated priority list and add four contrast papers on early effect evaluation, user adoption, integrative framing and AI risk governance. They are reading prompts, not replacements for the papers. `not_reported` is preserved rather than filled with assumptions.

{"".join(sections)}
"""


def build_direction_options() -> str:
    return """# Research Direction Options

This document presents alternative research paths rather than one recommended topic. Each option is supported by corpus IDs and should be narrowed only after checking field access, data availability, time horizon, and theoretical interest.

## Option 1: Post-occupancy evaluation of learning spaces

**Research question.** After a space redesign, how do participation, collaboration, sustained use and learning outcomes change, and which spatial features plausibly explain the change? [FLC-AF43D204976D] [FLC-B21262428AA9] [FLC-40CAC33A5980]

**Potential contribution.** Move from design description to outcome evaluation; distinguish satisfaction, observed behaviour and learning outcomes. [FLC-AF43D204976D] [FLC-B21262428AA9]

**Possible design.** Pre/post survey plus occupancy or booking data; comparison space or phased rollout; course-level qualitative follow-up. [FLC-AF43D204976D] [FLC-40CAC33A5980]

**Possible difficulty.** Student and course self-selection, lack of baseline data, and simultaneous teaching changes complicate attribution. [FLC-AF43D204976D]

## Option 2: Cross-institutional user-needs portfolio

**Research question.** Are mandatory, expected and attractive service needs stable across institution types and user groups, and how do they change after service experience? [FLC-D2B84F8B5474] [FLC-0A78AE1EC059] [FLC-E86EDD77A7D8]

**Potential contribution.** Test the external validity and temporal stability of KANO-based service priorities. [FLC-D2B84F8B5474] [FLC-E86EDD77A7D8]

**Possible design.** Multi-site repeated survey with measurement-invariance checks, linked where possible to actual service-use records. [FLC-D2B84F8B5474] [FLC-0A78AE1EC059]

**Possible difficulty.** Questionnaire equivalence and uneven implementation maturity across institutions. [FLC-D2B84F8B5474] [FLC-E86EDD77A7D8]

## Option 3: AI-Agent service trial with governance outcomes

**Research question.** Does an AI-Agent improve response quality, learning support or resource discovery relative to existing service, and at what cost in error, privacy and responsibility? [FLC-7E32D0AFAA75] [FLC-8177904692F9] [FLC-FADDE187EAA5]

**Potential contribution.** Connect technical performance, user outcomes and institutional governance in one evaluative design. [FLC-8177904692F9] [FLC-FADDE187EAA5]

**Possible design.** Bounded pilot with blind quality scoring, user task completion, escalation rates, privacy review and librarian workload measures. [FLC-7E32D0AFAA75] [FLC-FADDE187EAA5]

**Possible difficulty.** Rapid model change, data authorization, safety monitoring and reproducibility. [FLC-FADDE187EAA5]

## Option 4: Governance models for library transformation

**Research question.** How do library-led, jointly governed and institution-wide models differ in decision rights, budgets, service integration and accountability? [FLC-079C3D09D5B5] [FLC-5219CF870983] [FLC-CFCDAA38CDA0]

**Potential contribution.** Convert “top-level design” into observable governance variables and comparable organizational mechanisms. [FLC-5219CF870983] [FLC-CFCDAA38CDA0]

**Possible design.** Comparative case study with process tracing, governance documents, interviews and a common coding framework. [FLC-079C3D09D5B5] [FLC-CFCDAA38CDA0]

**Possible difficulty.** Sensitive organizational evidence and low comparability across institutions. [FLC-079C3D09D5B5]

## Option 5: Librarian role and learning-support effectiveness

**Research question.** How are learning planning, teaching-design consultation and data support embedded in courses, and how do librarian capabilities affect service outcomes? [FLC-54FCF9633BD1] [FLC-4979619F5099] [FLC-079C3D09D5B5]

**Potential contribution.** Link professional-role change to tasks, capabilities, workload and learner-facing outcomes. [FLC-54FCF9633BD1] [FLC-4979619F5099]

**Possible design.** Service-process mapping, competency assessment, course-embedded case studies and outcome tracking. [FLC-54FCF9633BD1] [FLC-079C3D09D5B5]

**Possible difficulty.** Multi-department delivery makes the library's independent contribution difficult to isolate. [FLC-079C3D09D5B5]

## Option 6: Comparative study across institution types

**Research question.** How do constraints and service priorities differ among research universities, applied undergraduate institutions, vocational colleges and discipline-specific settings? [FLC-E86EDD77A7D8] [FLC-309D35433AE0] [FLC-0970AAAB7995] [FLC-AF0DA0F0CDAE]

**Potential contribution.** Build a contextual typology rather than treating resource-rich universities as the universal model. [FLC-E86EDD77A7D8] [FLC-309D35433AE0]

**Possible design.** Maximum-variation case sampling with a shared minimum dataset and context-specific modules. [FLC-E86EDD77A7D8] [FLC-AF0DA0F0CDAE]

**Possible difficulty.** Institutional heterogeneity and measurement equivalence. [FLC-E86EDD77A7D8] [FLC-0970AAAB7995]

## Option 7: Conceptual clarification and measurement model

**Research question.** What are the necessary conditions, boundaries and outcome variables when “future learning center” refers to an organization, space, service system or policy project? [FLC-BFA71761FFA6] [FLC-39D9F87A8F5B] [FLC-77F19452C268]

**Potential contribution.** Provide a multi-level construct model that makes later empirical studies comparable. [FLC-BFA71761FFA6] [FLC-77F19452C268]

**Possible design.** Concept analysis, expert elicitation, item generation and pilot construct validation. [FLC-BFA71761FFA6] [FLC-39D9F87A8F5B]

**Possible difficulty.** A single definition may erase legitimate contextual variation; the model should preserve multiple levels. [FLC-BFA71761FFA6]

## Option 8: International comparison with controlled context

**Research question.** Which spatial, service and governance mechanisms remain comparable after accounting for institutional mission, finance, user structure and policy environment? [FLC-1C677BADE9E4] [FLC-5A559DBBB63F] [FLC-1FBEE4B5422D]

**Potential contribution.** Distinguish transferable mechanisms from context-dependent practices. [FLC-5A559DBBB63F] [FLC-1FBEE4B5422D]

**Possible design.** Matched-institution comparison using common indicators, documents, interviews and service-use evidence. [FLC-1C677BADE9E4] [FLC-5A559DBBB63F]

**Possible difficulty.** Source-language access, data comparability and incomplete public information. [FLC-5A559DBBB63F] [FLC-1FBEE4B5422D]

## How to choose without prematurely fixing a topic

Use four filters:

1. **Unit of analysis:** organization, space, service, technology or user.
2. **Evidence access:** documents only, survey participants, behavioural logs, institutional records or intervention access.
3. **Claim ambition:** descriptive, explanatory, evaluative or causal.
4. **Time horizon:** cross-sectional, repeated, longitudinal or implementation study.

A viable direction is the intersection of a bounded unit, accessible evidence and a claim type the design can support. The corpus does not justify selecting one option for every researcher. [FLC-BFA71761FFA6] [FLC-D2B84F8B5474] [FLC-AF43D204976D] [FLC-CFCDAA38CDA0]
"""


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Research reader  •  ")
    set_run_font(run, size=9, color="667085")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def style_callout(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F6F9")
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "2E74B5")
    borders.append(left)
    p_pr.append(borders)
    paragraph.paragraph_format.left_indent = Inches(0.14)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.2


def configure_doc(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("FUTURE LEARNING CENTER LITERATURE READER")
    set_run_font(run, size=8.5, color="667085", bold=True)
    add_page_number(section.footer.paragraphs[0])


def add_rich_text(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=9.5, color="7A5A00")
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def add_markdown_table(doc, lines):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [rows[0]] + rows[2:]
    cols = len(rows[0])
    widths = {
        2: [2700, 6660],
        3: [4200, 1800, 3360],
        4: [1600, 2300, 2500, 2960],
    }.get(cols, [9360 // cols] * cols)
    widths[-1] += 9360 - sum(widths)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            if c_idx > 0 and re.fullmatch(r"[\d,]+", value):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich_text(p, value)
            for run in p.runs:
                set_run_font(run, size=8.5 if cols >= 4 else 9)
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string("0B2545")
            if r_idx == 0:
                shade_cell(cell, "E8EEF5")
    set_table_geometry(table, widths)
    for row in table.rows:
        row._tr.get_or_add_trPr()
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def markdown_to_docx(markdown: str, output: Path):
    doc = Document()
    configure_doc(doc)
    lines = markdown.splitlines()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(70)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("Future Learning Center\nLiterature Reader")
    set_run_font(run, size=28, color="203748", bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(58)
    run = subtitle.add_run("A researcher-oriented notebook for interpreting the corpus and framing future inquiry")
    set_run_font(run, size=12, color="667085", italic=True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Evidence base: 215 PDF records  •  synthesis from validated database outputs")
    set_run_font(run, size=10, color="667085")
    doc.add_page_break()

    i = 1
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[-:| ]+\|$", lines[i + 1]):
            table_lines = [line, lines[i + 1]]
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("> "):
            p = doc.add_paragraph()
            style_callout(p)
            add_rich_text(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_rich_text(p, re.sub(r"^\d+\.\s+", "", line))
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_rich_text(p, line)
        i += 1

    core = doc.core_properties
    core.title = "Future Learning Center Literature Reader"
    core.subject = "Research-oriented literature reader"
    core.author = "Research synthesis"
    core.keywords = "future learning center, library transformation, research reader"
    doc.save(output)


def validate_ids(texts: dict[str, str], valid_ids: set[str]):
    invalid = {}
    for name, text in texts.items():
        cited = set(re.findall(r"FLC-[0-9A-F]{12}", text))
        bad = cited - valid_ids
        if bad:
            invalid[name] = sorted(bad)
    if invalid:
        raise RuntimeError(f"Invalid cited IDs: {invalid}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    args = parser.parse_args()
    root = args.root.resolve()
    outputs = {
        "reader": root / "future_learning_center_reader.md",
        "docx": root / "future_learning_center_reader.docx",
        "notes": root / "key_paper_notes.md",
        "directions": root / "research_direction_options.md",
    }
    existing = [str(path) for path in outputs.values() if path.exists()]
    if existing:
        raise FileExistsError(f"Refusing to overwrite: {existing}")

    records = read_records(root)
    reader = build_reader(records)
    notes = build_key_notes(records)
    directions = build_direction_options()
    validate_ids(
        {"reader": reader, "notes": notes, "directions": directions},
        {record["stable_id"] for record in records},
    )
    outputs["reader"].write_text(reader, encoding="utf-8", newline="\n")
    outputs["notes"].write_text(notes, encoding="utf-8", newline="\n")
    outputs["directions"].write_text(directions, encoding="utf-8", newline="\n")
    markdown_to_docx(reader, outputs["docx"])
    print(
        json.dumps(
            {
                "records": len(records),
                "reader_chars": len(reader),
                "key_notes": len(KEY_NOTES),
                "outputs": {key: str(value) for key, value in outputs.items()},
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
